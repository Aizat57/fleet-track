import csv
import io

from django.db.models import Max, Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import ComplianceRecord, Vehicle


def fleet_qs():
    return Vehicle.objects.exclude(status__in=Vehicle.INACTIVE_STATUSES)


def with_expiries(qs):
    return qs.annotate(
        ins_expiry=Max('compliance_records__expiry_date', filter=Q(compliance_records__kind=ComplianceRecord.Kind.INSURANCE)),
        rt_expiry=Max('compliance_records__expiry_date', filter=Q(compliance_records__kind=ComplianceRecord.Kind.ROAD_TAX)),
    )


def current_records(vehicles=None):
    """Latest-expiry record per (vehicle, kind) for in-fleet vehicles: {(vehicle_id, kind): record}."""
    qs = ComplianceRecord.objects.select_related('vehicle', 'responsible')
    qs = qs.filter(vehicle__in=vehicles) if vehicles is not None else qs.exclude(vehicle__status__in=Vehicle.INACTIVE_STATUSES)
    current = {}
    for r in qs.order_by('expiry_date', 'id'):
        current[(r.vehicle_id, r.kind)] = r
    return current


def tabular_response(filename, headers, rows, fmt='xlsx', title='FleetTrack'):
    """Stream rows as .xlsx (default) or .csv."""
    stamp = timezone.localdate().strftime('%Y%m%d')
    if fmt == 'csv':
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        writer.writerows(rows)
        resp = HttpResponse('﻿' + buf.getvalue(), content_type='text/csv; charset=utf-8')
        resp['Content-Disposition'] = f'attachment; filename="{filename}-{stamp}.csv"'
        return resp
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
    for row in rows:
        ws.append(list(row))
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 48)
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{filename}-{stamp}.xlsx"'
    return resp
