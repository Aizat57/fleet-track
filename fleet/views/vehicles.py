import csv
import io
from datetime import date, datetime, timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import QueryDict
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST
from openpyxl import load_workbook

from .. import finance
from ..forms import (VEHICLE_BASIC, VEHICLE_OPERATIONS, VEHICLE_OWNERSHIP, VEHICLE_STEPS, AssignmentForm,
                     DocumentForm, ImportForm, ServiceTaskForm, InitialComplianceForm, VehicleForm, VehicleSetupForm)
from ..models import (Assignment, ComplianceRecord, Department, Document, Driver, Expense, Location, Vehicle,
                      VehicleDraft)
from ..utils import log_activity, safe_next, paginate, require_perm
from .common import fleet_qs, tabular_response, with_expiries

EXPIRY_FILTERS = [('valid', 'Valid'), ('expiring', 'Expiring ≤ 30 days'), ('expired', 'Expired'),
                  ('missing', 'Not recorded')]
SORTS = {'reg': 'registration_no', 'type': 'vehicle_type', 'status': 'status', 'ins': 'ins_expiry', 'rt': 'rt_expiry',
         'make': 'make'}


def _expiry_filter(qs, field, value, today, applies_to=None):
    soon = today + timedelta(days=30)
    if value == 'expired':
        return qs.filter(**{f'{field}__lt': today})
    if value == 'expiring':
        return qs.filter(**{f'{field}__gte': today, f'{field}__lte': soon})
    if value == 'valid':
        return qs.filter(**{f'{field}__gt': soon})
    if value == 'missing':
        qs = qs.filter(**{f'{field}__isnull': True})
        return qs.filter(vehicle_type__in=applies_to) if applies_to is not None else qs
    return qs


def filtered_vehicles(params):
    today = timezone.localdate()
    status = params.get('status', '')
    qs = with_expiries(Vehicle.objects.select_related('department', 'location'))
    if status == 'all':
        pass
    elif status:
        qs = qs.filter(status=status)
    else:
        qs = qs.exclude(status__in=Vehicle.INACTIVE_STATUSES)
    if q := params.get('q', '').strip():
        qs = qs.filter(Q(registration_no__icontains=q) | Q(make__icontains=q) | Q(model__icontains=q)
                       | Q(chassis_no__icontains=q) | Q(assignments__driver__name__icontains=q,
                                                       assignments__end_date__isnull=True)).distinct()
    for key, field in (('type', 'vehicle_type'), ('ownership', 'ownership_type'), ('department', 'department_id'),
                       ('location', 'location_id')):
        if params.get(key):
            qs = qs.filter(**{field: params[key]})
    insured = [t for t in Vehicle.Type.values if t not in Vehicle.INSURANCE_EXEMPT_TYPES]
    qs = _expiry_filter(qs, 'ins_expiry', params.get('insurance'), today, insured)
    qs = _expiry_filter(qs, 'rt_expiry', params.get('road_tax'), today, Vehicle.ROAD_TAX_TYPES)
    if params.get('expiring'):
        soon = today + timedelta(days=30)
        qs = qs.filter(Q(ins_expiry__lte=soon) | Q(rt_expiry__lte=soon))
    sort = params.get('sort', 'reg')
    field = SORTS.get(sort.lstrip('-'), 'registration_no')
    qs = qs.order_by(('-' if sort.startswith('-') else '') + field, 'registration_no')
    return qs.prefetch_related('assignments__driver')


@login_required
def vehicle_list(request):
    qs = filtered_vehicles(request.GET)
    if fmt := request.GET.get('export'):
        return export_vehicles(qs, fmt)
    today = timezone.localdate()
    fleet = with_expiries(fleet_qs())
    soon = today + timedelta(days=30)
    summary = fleet.aggregate(total=Count('id'), active=Count('id', filter=Q(status=Vehicle.Status.ACTIVE)),
                              maintenance=Count('id', filter=Q(status=Vehicle.Status.MAINTENANCE)))
    summary['expiring'] = fleet.filter(Q(ins_expiry__lte=soon) | Q(rt_expiry__lte=soon)).count()
    filter_keys = ('q', 'type', 'status', 'ownership', 'insurance', 'road_tax', 'department', 'location', 'expiring')
    active_filters = sum(1 for k in filter_keys[1:] if request.GET.get(k))
    return render(request, 'fleet/vehicles/list.html', {
        'page': paginate(request, qs, 25),
        'summary': summary,
        'type_groups': [(g, [(t, Vehicle.Type(t).label) for t in types]) for g, types in Vehicle.TYPE_GROUPS],
        'statuses': Vehicle.Status.choices,
        'ownerships': Vehicle.Ownership.choices,
        'expiry_filters': EXPIRY_FILTERS,
        'departments': Department.objects.all(),
        'locations': Location.objects.all(),
        'active_filters': active_filters,
        'has_any': Vehicle.objects.exists(),
        'drafts': VehicleDraft.objects.filter(user=request.user)[:5],
        'sort': request.GET.get('sort', 'reg'),
        'active_nav': 'vehicles',
    })


def export_vehicles(qs, fmt):
    headers = ['Registration No.', 'Type', 'Make', 'Model', 'Year', 'Status', 'Ownership', 'Primary Driver',
               'Department', 'Location', 'Insurance Expiry', 'Road Tax Expiry', 'Chassis No.', 'Engine No.',
               'Odometer (km)']
    rows = [[v.registration_no, v.get_vehicle_type_display(), v.make, v.model, v.year, v.get_status_display(),
             v.get_ownership_type_display(), str(v.primary_driver or ''), str(v.department or ''),
             str(v.location or ''), v.ins_expiry, v.rt_expiry, v.chassis_no, v.engine_no, v.odometer_km] for v in qs]
    return tabular_response('vehicles', headers, rows, fmt, 'Vehicles')


# ---------------------------------------------------------------- add vehicle wizard

def _draft_initial(data, fields, prefix=None):
    initial = {}
    for name in fields:
        key = f'{prefix}-{name}' if prefix else name
        if key in data:
            values = data[key]
            initial[name] = values if len(values) > 1 else values[0]
    return initial


def _wizard_forms(request, draft=None):
    if request.method == 'POST' and request.POST.get('action') != 'draft':
        args = (request.POST, request.FILES)
        return (VehicleForm(*args), InitialComplianceForm(*args, prefix='ins', kind=ComplianceRecord.Kind.INSURANCE),
                InitialComplianceForm(*args, prefix='rt', kind=ComplianceRecord.Kind.ROAD_TAX), VehicleSetupForm(*args))
    data = draft.data if draft else {}
    return (
        VehicleForm(initial=_draft_initial(data, VehicleForm.Meta.fields)),
        InitialComplianceForm(prefix='ins', kind=ComplianceRecord.Kind.INSURANCE,
                              initial=_draft_initial(data, ['provider', 'reference_no', 'coverage', 'start_date',
                                                            'expiry_date', 'cost'], 'ins')),
        InitialComplianceForm(prefix='rt', kind=ComplianceRecord.Kind.ROAD_TAX,
                              initial=_draft_initial(data, ['provider', 'reference_no', 'start_date', 'expiry_date',
                                                            'cost'], 'rt')),
        VehicleSetupForm(initial=_draft_initial(data, ['driver'])),
    )


def _save_document(request, upload, category, vehicle, name=None, **links):
    return Document.objects.create(
        name=name or upload.name.rsplit('.', 1)[0][:160], category=category, file=upload,
        original_filename=upload.name[:255], content_type=getattr(upload, 'content_type', '') or '',
        size=upload.size, vehicle=vehicle, uploaded_by=request.user, **links)


def create_compliance(request, vehicle, kind, data, payment_status=Expense.Payment.PAID, record_expense=True):
    previous = vehicle.compliance_records.filter(kind=kind)
    record = ComplianceRecord.objects.create(
        kind=kind, vehicle=vehicle, provider=data.get('provider') or '', reference_no=data.get('reference_no') or '',
        coverage=data.get('coverage') or '', sum_insured=data.get('sum_insured'), start_date=data.get('start_date'),
        expiry_date=data['expiry_date'], cost=data.get('cost'), responsible=data.get('responsible'),
        workflow=data.get('workflow') or '', is_renewal=previous.exists(), created_by=request.user)
    previous.exclude(pk=record.pk).update(workflow='')
    if record.cost and record_expense:
        label = record.get_kind_display()
        Expense.objects.create(
            vehicle=vehicle, date=record.start_date or timezone.localdate(),
            category=Expense.Category.INSURANCE if kind == ComplianceRecord.Kind.INSURANCE else Expense.Category.ROAD_TAX,
            description=f'{label} {record.start_date or ""}–{record.expiry_date}'.replace(' –', ' to ', 1).strip(),
            supplier=record.provider, invoice_no=record.reference_no, amount=record.cost,
            payment_status=payment_status, compliance_record=record, created_by=request.user)
    return record


@require_perm('fleet.add_vehicle')
def vehicle_create(request):
    draft = None
    if draft_id := request.GET.get('draft') or request.POST.get('draft_id'):
        draft = get_object_or_404(VehicleDraft, pk=draft_id, user=request.user)

    if request.method == 'POST' and request.POST.get('action') == 'draft':
        data = {k: v for k, v in request.POST.lists() if k not in ('csrfmiddlewaretoken', 'action', 'draft_id')}
        draft = draft or VehicleDraft(user=request.user)
        draft.data = data
        draft.step = int(request.POST.get('step') or 1)
        draft.save()
        messages.success(request, 'Draft saved. Attached files are not kept in drafts; add them when you finish.')
        return redirect('fleet:vehicle_list')

    vform, ins_form, rt_form, setup = _wizard_forms(request, draft)
    if request.method == 'POST':
        valid = all([vform.is_valid(), ins_form.is_valid(), rt_form.is_valid(), setup.is_valid()])
        if valid:
            with transaction.atomic():
                vehicle = vform.save(commit=False)
                vehicle.created_by = request.user
                vehicle.save()
                for form in (ins_form, rt_form):
                    if form.has_details() and vehicle.requires(form.kind):
                        create_compliance(request, vehicle, form.kind, form.cleaned_data)
                driver = setup.cleaned_data.get('driver')
                if driver and vehicle.takes_driver:
                    Assignment.objects.create(vehicle=vehicle, driver=driver, is_primary=True)
                for key, category in (('registration_file', Document.Category.REGISTRATION),
                                      ('insurance_file', Document.Category.INSURANCE),
                                      ('road_tax_file', Document.Category.ROAD_TAX)):
                    if upload := setup.cleaned_data.get(key):
                        _save_document(request, upload, category, vehicle,
                                       name=f'{Document.Category(category).label} – {vehicle.registration_no}')
                for upload in setup.cleaned_data.get('other_files') or []:
                    _save_document(request, upload, Document.Category.OTHER, vehicle)
                if draft:
                    draft.delete()
                log_activity(request, 'vehicle_added', f'Vehicle {vehicle.registration_no} added', vehicle,
                             vehicle.get_absolute_url())
            messages.success(request, f'{vehicle.registration_no} has been added to the fleet.')
            return redirect(vehicle)
        messages.error(request, 'Some details need attention. Steps with problems are marked in red.')

    step_fields = {
        'basic': [vform[f] for f in VEHICLE_BASIC],
        'ownership': [vform[f] for f in VEHICLE_OWNERSHIP],
        'operations': [vform[f] for f in VEHICLE_OPERATIONS],
    }
    step_errors = {
        'basic': any(vform[f].errors for f in VEHICLE_BASIC),
        'ownership': any(vform[f].errors for f in VEHICLE_OWNERSHIP),
        'compliance': bool(ins_form.errors or rt_form.errors),
        'driver': bool(setup['driver'].errors),
        'operations': any(vform[f].errors for f in VEHICLE_OPERATIONS),
        'review': any(setup[f].errors for f in ('registration_file', 'insurance_file', 'road_tax_file', 'other_files')),
    }
    return render(request, 'fleet/vehicles/wizard.html', {
        'vform': vform, 'ins_form': ins_form, 'rt_form': rt_form, 'setup': setup, 'draft': draft,
        'steps': VEHICLE_STEPS, 'step_fields': step_fields, 'step_errors': step_errors,
        'start_step': (draft.step if draft and request.method == 'GET' else 1),
        'has_drivers': Driver.objects.exclude(status=Driver.Status.INACTIVE).exists(),
        'insured_types': ','.join(t for t in Vehicle.Type.values if t not in Vehicle.INSURANCE_EXEMPT_TYPES),
        'uninsured_types': ','.join(Vehicle.INSURANCE_EXEMPT_TYPES),
        'road_tax_types': ','.join(Vehicle.ROAD_TAX_TYPES),
        'no_road_tax_types': ','.join(t for t in Vehicle.Type.values if t not in Vehicle.ROAD_TAX_TYPES),
        'active_nav': 'vehicles',
    })


@require_POST
@login_required
def draft_delete(request, pk):
    get_object_or_404(VehicleDraft, pk=pk, user=request.user).delete()
    messages.success(request, 'Draft discarded.')
    return redirect('fleet:vehicle_list')


@require_perm('fleet.change_vehicle')
def vehicle_edit(request, pk):
    vehicle = get_object_or_404(Vehicle, pk=pk)
    form = VehicleForm(request.POST or None, request.FILES or None, instance=vehicle)
    if request.method == 'POST' and form.is_valid():
        old_status = Vehicle.objects.get(pk=pk).get_status_display()
        vehicle = form.save()
        if old_status != vehicle.get_status_display():
            log_activity(request, 'status_changed',
                         f'{vehicle.registration_no} status changed from {old_status} to {vehicle.get_status_display()}',
                         vehicle, vehicle.get_absolute_url())
        else:
            log_activity(request, 'vehicle_updated', f'{vehicle.registration_no} details updated', vehicle,
                         vehicle.get_absolute_url())
        messages.success(request, 'Vehicle details saved.')
        return redirect(vehicle)
    sections = [('Basic Information', VEHICLE_BASIC), ('Ownership', VEHICLE_OWNERSHIP),
                ('Operational Information', VEHICLE_OPERATIONS)]
    return render(request, 'fleet/vehicles/edit.html', {
        'vehicle': vehicle, 'form': form,
        'sections': [(title, [form[f] for f in fields]) for title, fields in sections],
        'active_nav': 'vehicles',
    })


# ---------------------------------------------------------------- vehicle profile

VEHICLE_TABS = [('overview', 'Overview'), ('ownership', 'Ownership'), ('insurance', 'Insurance'),
                ('road_tax', 'Road Tax'), ('drivers', 'Drivers'), ('maintenance', 'Maintenance & Repairs'),
                ('expenses', 'Expenses'), ('documents', 'Documents'), ('history', 'History')]


def _history_entries(vehicle):
    entries = []
    for rec in vehicle.history.select_related('history_user')[:40]:
        changes = []
        prev = rec.prev_record
        if prev is not None:
            delta = rec.diff_against(prev)
            for c in delta.changes:
                f = Vehicle._meta.get_field(c.field)
                changes.append({'field': f.verbose_name.capitalize(), 'old': c.old, 'new': c.new})
        kind = {'+': 'Created', '~': 'Updated', '-': 'Deleted'}[rec.history_type]
        if kind == 'Updated' and not changes:
            continue
        entries.append({'when': rec.history_date, 'user': rec.history_user, 'kind': kind, 'changes': changes})
    return entries


@login_required
def vehicle_detail(request, pk):
    vehicle = get_object_or_404(
        Vehicle.objects.select_related('department', 'location', 'created_by')
        .prefetch_related('compliance_records__responsible', 'assignments__driver', 'documents'), pk=pk)
    today = timezone.localdate()
    month_start = today.replace(day=1)
    expenses = vehicle.expenses.all()
    month_spend = finance.summary(month_start, today, expenses)
    records = sorted(vehicle.compliance_records.all(), key=lambda r: r.expiry_date, reverse=True)
    tab = request.GET.get('tab', 'overview')
    if tab not in dict(VEHICLE_TABS):
        tab = 'overview'
    return render(request, 'fleet/vehicles/detail.html', {
        'vehicle': vehicle,
        'tabs': VEHICLE_TABS,
        'tab': tab,
        'insurance': vehicle.current_insurance,
        'road_tax': vehicle.current_road_tax,
        'insurance_records': [r for r in records if r.kind == ComplianceRecord.Kind.INSURANCE],
        'road_tax_records': [r for r in records if r.kind == ComplianceRecord.Kind.ROAD_TAX],
        'assignments': vehicle.assignments.all(),
        'active_assignments': vehicle.active_assignments,
        'assignment_form': AssignmentForm(),
        'upload_form': DocumentForm(initial={'vehicle': vehicle.pk}),
        'task_form': ServiceTaskForm(initial={'vehicle': vehicle.pk}),
        'breakdowns': vehicle.breakdowns.all()[:50],
        'service_tasks': vehicle.service_tasks.all()[:50],
        'expenses': expenses.select_related('created_by')[:50],
        'expense_count': expenses.count(),
        'month_spend': month_spend,
        'has_expenses': expenses.exists(),
        'lifetime_spend': finance.actual(expenses).aggregate(s=Sum('amount'))['s'],
        'documents': vehicle.documents.select_related('uploaded_by'),
        'history': _history_entries(vehicle),
        'activities': vehicle.activities.select_related('actor')[:30],
        'today': today,
        'odometer_types': Vehicle.ODOMETER_TYPES,
        'hours_types': Vehicle.ENGINE_HOURS_TYPES,
        'active_nav': 'vehicles',
    })


@require_POST
@require_perm('fleet.add_assignment')
def assignment_create(request, pk):
    vehicle = get_object_or_404(Vehicle, pk=pk)
    form = AssignmentForm(request.POST)
    if not vehicle.takes_driver:
        messages.error(request, 'Trailers are not assigned drivers; assign the prime mover instead.')
    elif form.is_valid():
        a = form.save(commit=False)
        a.vehicle = vehicle
        with transaction.atomic():
            if a.is_primary:
                for old in vehicle.assignments.filter(end_date__isnull=True, is_primary=True):
                    old.end_date = a.start_date
                    old.save()
            if vehicle.assignments.filter(driver=a.driver, end_date__isnull=True).exists():
                messages.info(request, f'{a.driver} is already assigned to this vehicle.')
                return redirect(f'{vehicle.get_absolute_url()}?tab=drivers')
            a.save()
        log_activity(request, 'driver_assigned', f'{a.driver} assigned to {vehicle.registration_no}'
                     + (' as primary driver' if a.is_primary else ''), vehicle, vehicle.get_absolute_url() + '?tab=drivers')
        messages.success(request, f'{a.driver} assigned.')
    else:
        messages.error(request, 'Choose a driver to assign.')
    return redirect(f'{vehicle.get_absolute_url()}?tab=drivers')


@require_POST
@require_perm('fleet.change_assignment')
def assignment_end(request, pk):
    a = get_object_or_404(Assignment.objects.select_related('vehicle', 'driver'), pk=pk)
    a.end_date = timezone.localdate()
    a.save()
    log_activity(request, 'driver_unassigned', f'{a.driver} unassigned from {a.vehicle.registration_no}', a.vehicle,
                 a.vehicle.get_absolute_url() + '?tab=drivers')
    messages.success(request, f'{a.driver} is no longer assigned to {a.vehicle}.')
    return redirect(safe_next(request, f'{a.vehicle.get_absolute_url()}?tab=drivers'))


@require_POST
@require_perm('fleet.change_vehicle')
def vehicle_photo(request, pk):
    vehicle = get_object_or_404(Vehicle, pk=pk)
    upload = request.FILES.get('image')
    if upload and upload.content_type.startswith('image/') and upload.size <= 10 * 1024 * 1024:
        vehicle.image = upload
        vehicle.save(update_fields=['image', 'updated_at'])
        messages.success(request, 'Photo updated.')
    else:
        messages.error(request, 'Choose an image file under 10 MB.')
    return redirect(vehicle)


# ---------------------------------------------------------------- import

IMPORT_COLUMNS = ['registration_no', 'vehicle_type', 'make', 'model', 'year', 'colour', 'chassis_no', 'engine_no',
                  'ownership_type', 'registered_owner', 'status', 'department', 'location', 'fuel_type', 'odometer_km',
                  'insurance_expiry', 'road_tax_expiry']


TYPE_ALIASES = {
    'kereta': 'car', 'motosikal': 'motorcycle', 'motor': 'motorcycle', 'basikal': 'bicycle', 'lori': 'truck',
    'lori tangki': 'tanker', 'treler': 'trailer', 'kepala treler': 'prime_mover', 'bas': 'bus', 'pikap': 'pickup',
    'jentolak': 'excavator', 'jengkaut': 'excavator', 'kren': 'crane', 'traktor': 'tractor', 'bot': 'boat',
    'kapal': 'boat', 'helikopter': 'helicopter', 'heli': 'helicopter', 'kapal terbang': 'aircraft',
    'pesawat': 'aircraft', 'dron': 'drone', 'lain-lain': 'other',
}


def _choice_lookup(choices):
    lookup = {}
    for value, label in choices:
        lookup[value.lower()] = value
        lookup[label.lower()] = value
        for part in label.lower().replace('/', ' / ').split(' / '):
            lookup.setdefault(part.strip(), value)
    return lookup


def _read_rows(upload):
    if upload.name.lower().endswith('.csv'):
        text = upload.read().decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    else:
        wb = load_workbook(upload, read_only=True, data_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    if not rows:
        return []
    headers = [str(h or '').strip().lower().replace(' ', '_').replace('.', '') for h in rows[0]]
    aliases = {'registration_number': 'registration_no', 'reg_no': 'registration_no', 'type': 'vehicle_type',
               'ownership': 'ownership_type', 'odometer': 'odometer_km'}
    headers = [aliases.get(h, h) for h in headers]
    return [dict(zip(headers, r)) for r in rows[1:] if any(c not in (None, '') for c in r)]


def _to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value or '').strip()
    if not s:
        return None
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return parse_date(s)


@require_perm('fleet.add_vehicle')
def vehicle_import(request):
    form = ImportForm(request.POST or None, request.FILES or None)
    results = None
    if request.GET.get('template'):
        example = ['WXY 1234', 'Van', 'Toyota', 'Hiace', 2021, 'White', 'KDH200-123456', '2KD-987654', 'Company Owned',
                   'Syarikat Contoh Sdn Bhd', 'Active', 'Logistics', 'Shah Alam', 'Diesel', 85000, '2027-03-31',
                   '2027-03-31']
        return tabular_response('vehicle-import-template', IMPORT_COLUMNS, [example], 'xlsx', 'Vehicles')
    if request.method == 'POST' and form.is_valid():
        try:
            rows = _read_rows(form.cleaned_data['file'])
        except Exception:
            rows = None
            form.add_error('file', 'The file could not be read. Use the template and save as .xlsx or .csv.')
        if rows is not None:
            results = {'created': [], 'errors': []}
            lookups = {'vehicle_type': {**_choice_lookup(Vehicle.Type.choices), **TYPE_ALIASES},
                       'ownership_type': _choice_lookup(Vehicle.Ownership.choices),
                       'status': _choice_lookup(Vehicle.Status.choices),
                       'fuel_type': _choice_lookup(Vehicle.Fuel.choices)}
            for n, row in enumerate(rows, start=2):
                data = QueryDict(mutable=True)
                for col in VehicleForm.Meta.fields:
                    val = row.get(col)
                    if val in (None, ''):
                        continue
                    val = str(val).strip()
                    if col in lookups:
                        val = lookups[col].get(val.lower(), val)
                    if col in ('department', 'location'):
                        model = Department if col == 'department' else Location
                        val = model.objects.get_or_create(name=val)[0].pk
                    data[col] = val
                data.setdefault('ownership_type', Vehicle.Ownership.OWNED)
                data.setdefault('status', Vehicle.Status.ACTIVE)
                vform = VehicleForm(data)
                ins, rt = _to_date(row.get('insurance_expiry')), _to_date(row.get('road_tax_expiry'))
                if vform.is_valid():
                    with transaction.atomic():
                        v = vform.save(commit=False)
                        v.created_by = request.user
                        v.save()
                        if ins:
                            create_compliance(request, v, ComplianceRecord.Kind.INSURANCE, {'expiry_date': ins})
                        if rt:
                            create_compliance(request, v, ComplianceRecord.Kind.ROAD_TAX, {'expiry_date': rt})
                    results['created'].append(v)
                else:
                    msgs = [f'{vform.fields[k].label if k in vform.fields else k}: {" ".join(e)}'
                            for k, e in vform.errors.items()]
                    results['errors'].append({'row': n, 'reg': row.get('registration_no') or '—', 'messages': msgs})
            if results['created']:
                log_activity(request, 'vehicles_imported', f'{len(results["created"])} vehicles imported')
                messages.success(request, f'{len(results["created"])} vehicle(s) imported.')
    return render(request, 'fleet/vehicles/import.html', {'form': form, 'results': results,
                                                          'columns': IMPORT_COLUMNS, 'active_nav': 'vehicles'})
